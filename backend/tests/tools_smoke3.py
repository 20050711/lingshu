"""4.1 工具加强冒烟测试：file_parse 三层分流 / chart_builder 设计常量 /
doc_export 美化与 pptx 图表内嵌 / html_report / officecli_adapter 预览。

2026-09-17：sql_query stats 一节随数据查询线下线删除。

用法: python tests/tools_smoke3.py
"""
import asyncio
import sys
import uuid
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.agent.tools import ToolContext, get_tool
from app.services.chart_builder import build_echarts_option

CTX = ToolContext("smoke3", 1, "demo", "employee", "smoke", "/data/outputs/smoke3/1", user_id=1)


def _mk_xlsx(path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["城市", "消费", "展现量"])
    for c, v, s in [("北京", 32.5, 3200), ("上海", 28.9, 2850), ("广州", 12.3, 1200)]:
        ws.append([c, v, s])
    wb.save(str(path))


def _mk_pdf(path: Path) -> None:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "PDF Smoke Test 12345", fontsize=16)
    doc.save(str(path))
    doc.close()


async def main() -> int:
    results = []
    # 测试文件放在产出目录内（validate_readable_path 仅放行当前会话上传/产出目录）
    out_root = Path("/data/outputs/smoke3/1")
    out_root.mkdir(parents=True, exist_ok=True)

    # 1. file_parse 三层分流：xlsx → excel-parser 深度解析（stats + chunks）
    xlsx = out_root / "smoke3_data.xlsx"
    _mk_xlsx(xlsx)
    t = get_tool("file_parse")
    r = await t.handler({"file_path": str(xlsx)}, CTX)
    ok = "sheets" in r and r["sheets"][0].get("stats", {}).get("消费") and "chunks" in r
    results.append(("file_parse xlsx 深度解析(stats+chunks)", ok,
                    f"stats={r['sheets'][0]['stats'].get('消费') if 'sheets' in r else r.get('error', '')[:40]} chunks={'chunks' in r}"))

    # 2. file_parse 三层分流：pdf → markitdown
    pdf = out_root / "smoke3_doc.pdf"
    _mk_pdf(pdf)
    r = await t.handler({"file_path": str(pdf)}, CTX)
    ok = "markdown" in r and "PDF Smoke Test" in r["markdown"]
    results.append(("file_parse pdf markitdown", ok, (r.get("markdown") or "")[:40].replace(chr(10), " ")))

    # 3. file_parse 越狱防护（保留）
    r = await t.handler({"file_path": "/etc/passwd"}, CTX)
    ok = "error" in r and "无权" in r["error"]
    results.append(("file_parse 越狱防护", ok, (r.get("error") or "")[:40]))

    # 4. sql_query stats → 2026-09-17 随数据查询线下线删除（stats 现由 _col_stats 单测覆盖）

    # 5. chart_builder 设计常量（配色/字体/圆角/dataZoom）
    opt = build_echarts_option("bar", "测试", ["城市", "消费"], [["北京", 32], ["上海", 28]], x_column="城市", y_columns=["消费"])
    ok = opt["color"][0] == "#1a56db" and opt["xAxis"]["axisLabel"]["color"] == "#64748b" \
        and opt["series"][0].get("barMaxWidth") == 40
    results.append(("chart_builder 设计常量", ok, f"palette={opt['color'][0]} barMaxWidth={opt['series'][0].get('barMaxWidth')}"))
    opt_line = build_echarts_option("line", "趋势", ["月", "值"], [[str(i), i] for i in range(30)], x_column="月", y_columns=["值"])
    ok = "dataZoom" in str(opt_line)
    results.append(("chart_builder 长序列 dataZoom", ok, ""))

    # 6. doc_export 四格式（pptx 带图表内嵌参数，渲染失败条目自动跳过不阻断）
    t = get_tool("doc_export")
    chart_id = str(uuid.uuid4())
    CTX.chart_registry = {chart_id: {"option": opt, "session_id": "smoke3"}}
    for fmt in ("docx", "xlsx", "pptx", "pdf"):
        r = await t.handler({"format": fmt, "title": "4.1 报告", "sections": [{"heading": "概览", "content": "消费增长 12%"}],
                             "table": {"headers": ["城市", "消费"], "rows": [["北京", 32.5]]},
                             "chart_ids": [chart_id]}, CTX)
        ok = "file_path" in r and Path("/data/outputs/smoke3/1", r["file_path"].split("/")[-1]).exists()
        results.append((f"doc_export {fmt} 生成", ok, (r.get("file_path") or r.get("error", ""))[-30:]))

    # 7. html_report（内联 echarts + 图表注入）
    t = get_tool("html_report")
    r = await t.handler({"title": "网页报告", "sections": [{"heading": "概览", "content": "数据如下"}],
                         "table": {"headers": ["城市", "消费"], "rows": [["北京", 32.5]]},
                         "chart_ids": [chart_id]}, CTX)
    if "file_path" in r:
        html = Path("/data/outputs/smoke3/1", r["file_path"].split("/")[-1]).read_text("utf-8")
        ok = "echarts.init" in html and "window.__CHARTS__" in html and "#1a56db" in html
        results.append(("html_report 内联图表", ok, f"{len(html) // 1024}KB"))
    else:
        results.append(("html_report 内联图表", False, str(r.get("error"))[:60]))

    # 8. officecli_adapter 预览（docx→HTML）
    from app.services.officecli_adapter import to_html

    docx = Path("/data/outputs/smoke3/1") / f"preview_{uuid.uuid4().hex[:4]}.docx"
    from docx import Document

    d = Document()
    d.add_heading("预览测试", 0)
    d.add_paragraph("内容段落")
    d.save(str(docx))
    try:
        html_path = await to_html(str(docx), "/data/outputs/smoke3/1")
        ok = html_path.endswith(".html") and Path(html_path).exists()
        results.append(("officecli_adapter docx→HTML", ok, Path(html_path).name))
    except RuntimeError as e:
        results.append(("officecli_adapter docx→HTML", False, str(e)[:60]))

    for name, ok, info in results:
        if ok is None:
            print(f"  - {name}: {info}")
            continue
        print(f"  {'✓' if ok else '✗'} {name}: {info}")
    checked = [ok for ok in (o for _, o, _ in results) if ok is not None]
    all_ok = all(checked)
    print(f"\ntools_smoke3: {'全部通过' if all_ok else f'{sum(1 for o in checked if not o)} 项失败'}")
    return 0 if all_ok else 1


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
