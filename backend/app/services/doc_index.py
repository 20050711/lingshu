"""通用文档索引层（2026-09-10）：解析分发 + 格式嗅探 + 切块。

知识库检索的通用索引层：切块、嗅探与候选预筛（检索工具与上传管线共用同一套实现）。

切块（chunk_document）：
- 先做**格式嗅探**：按内容结构判定（日期二级标题 + 说话人行），**不看文件名/扩展名**——
  换名字的导出文件照样识别；不命中就走通用段落策略（kb_chunker.chunk_text）
- 命中聊天记录导出 → 消息锚点策略：日期/整条消息为不可切断的边界，重叠回带完整消息
- 任何异常一律降级为"通用策略/不建块"，绝不抛给调用方（索引失败不能拖累上传）

解析分发（extract_text）四层逐层兜底：
- L1 专用解析器：txt/md/csv/docx/pptx/pdf/xlsx
- L2 LibreOffice 转换：doc/ppt/xls/wps/rtf/odt 等老 Office 家族（系统已装 soffice；单文件 30s 级，
  **仅供后台索引**，不进实时读取路径）
- L3 粗暴文本捞取：未知二进制按 UTF-16LE/UTF-8/GBK 试探 + 长串过滤（噪声多，标 quality=low）
- L4 取不到 → ok=False + reason（扫描件/加密/纯资源），调用方记 indexed=false，
  让 agent 知道"文件在但搜不到内容"，而不是以为"搜不到就是没有"
"""
from __future__ import annotations

import asyncio
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

from app.core.logging import get_logger
from app.core.text_utils import detect_binary_text, read_text_any_encoding
from app.services.kb_chunker import MAX_SIZE, MIN_SIZE, OVERLAP, chunk_text

logger = get_logger("services.doc_index")

# 纯文本类扩展（L1 直读，编码自检）
TEXT_EXTS = {".txt", ".md", ".csv"}
# 老 Office 家族（L2 LibreOffice 转换）
OFFICE_OLD_EXTS = {".doc", ".ppt", ".xls", ".wps", ".et", ".dps", ".rtf", ".odt", ".ods", ".odp"}
# 纯资源类（永远不会有块，只能按文件名/目录名找）
ASSET_EXTS = {
    ".psd", ".ai", ".eps", ".indd", ".sketch", ".fig", ".xd", ".ttf", ".otf", ".woff", ".woff2",
    ".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp", ".heic", ".tif", ".tiff", ".svg", ".raw",
    ".mp4", ".mov", ".avi", ".mkv", ".flv", ".wmv", ".mp3", ".wav", ".flac", ".aac", ".m4a",
    ".zip", ".rar", ".7z", ".tar", ".gz", ".iso", ".exe", ".dmg", ".apk",
    ".wxgf", ".silk", ".pcm", ".amr", ".dat", ".db", ".sqlite", ".bak", ".tmp",
}

# ---------- L1 专用解析器（自 kb_service 迁入，成为两库共用实现） ----------


XLSX_PREVIEW_MAX_CHARS = 20000  # xlsx_preview 拼接串的整体上限（老口径，勿动）
SHEET_MAX_CHARS = 20000         # 单 sheet 预览上限（一 sheet 一块时每块的上限）


def xlsx_preview_sheets(path: str) -> list[tuple[str, str]]:
    """逐 sheet 解析 xlsx：返回 [(sheet 名, 该 sheet 预览文本)]，**单 sheet 独立截断**。

    2026-09-18（问题 1）：原实现把所有 sheet 拼成一串再整体截断 20000 字，
    多 sheet 表格（实测导出 7 sheet）**靠后的 sheet 会被整段丢弃**
    ——不是被切碎，是在索引里根本不存在。改为逐 sheet 返回，调用方可一 sheet 一块。

    表格类**不切分**（结构数据切碎反而不可检索），整个 sheet 作为单块内容。
    """
    from app.agent.tools._col_stats import col_stats

    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out: list[tuple[str, str]] = []
    try:
        for ws in wb.worksheets:
            headers: list[str] = []
            keep: list[int] = []  # 非空表头列索引（max_col=500 会带出大量空列，展示前过滤）
            rows: list[list] = []
            # 2026-08-24（表格导出 dimension 异常）：read_only iter_rows 无参数只遍历 1 行 1 列
            # （同 import_pipeline 修复）——显式 max_row/max_col 强制全量
            for r_idx, row in enumerate(ws.iter_rows(values_only=True, max_row=1_000_000, max_col=500)):
                if row is None or all(v is None for v in row):
                    continue
                values = [str(v).strip() if v is not None else "" for v in row]
                if r_idx == 0:
                    headers = values
                    keep = [i for i, h in enumerate(headers) if h.strip()]
                    headers = [headers[i] for i in keep]
                    continue
                rows.append([values[i] for i in keep] if keep else values)
            total = len(rows)
            stats = col_stats(headers, rows)
            seg = rows[:50]
            lines = [f"## Sheet: {ws.title}（共 {total} 行数据）"]
            if headers:
                lines.append("表头: " + " | ".join(headers))
            if stats:
                lines.append("数值列统计: " + "；".join(
                    f"{h} sum={s['sum']} avg={s['avg']} max={s['max']} min={s['min']}"
                    for h, s in stats.items()
                ))
            if seg:
                lines.append("前 50 行预览:")
                for i, r in enumerate(seg, 1):
                    lines.append(f"{i}. " + " | ".join(r))
            out.append((ws.title, "\n".join(lines)[:SHEET_MAX_CHARS]))
    finally:
        wb.close()
    return out


def xlsx_preview(path: str) -> str:
    """整表预览文本（多 sheet 拼接后整体截断）——只吃字符串的调用方（知识库占位/LLM 摘要）沿用。"""
    return "\n\n".join(t for _title, t in xlsx_preview_sheets(path))[:XLSX_PREVIEW_MAX_CHARS]


def _markitdown_text(path: str) -> str:
    """markitdown 转 Markdown —— **与 file_parse 完全同源**（见 parse_document 注释）。"""
    from markitdown import MarkItDown

    return MarkItDown().convert(path).text_content or ""


def parse_document(path: str, ext: str) -> str:
    """L1 专用解析：txt/md/csv 编码自检直读；docx/pptx/pdf 走 **markitdown**。

    2026-09-10（根因修复）：索引与读取必须同源。原先索引走 python-docx/PyMuPDF、而 file_parse
    走 markitdown——同一个 pptx 两边文本量差 3 倍（6633 vs 20985 字），导致块的 char_start
    指向的是**另一个文本流**，agent 按 offset 续读会读到错误位置（比没有偏移更糟）。
    现统一为 markitdown（失败才回退专用解析器，此时文本不同源、偏移仅供近似定位）。
    """
    if ext in {".docx", ".pptx", ".pdf"}:
        try:
            md = _markitdown_text(path)
            if md.strip():
                return md
        except Exception as e:
            logger.warning("markitdown 解析失败，回退专用解析器 %s: %s", Path(path).name, str(e)[:120])
    if ext in TEXT_EXTS:
        # 2026-08-27（乱码修复）：编码检测——GBK/GB2312/UTF-16 文件原硬编码 UTF-8 解码乱码；
        # 二进制内容（zip 改名 .md 等）检测后明确报错，不入库乱码
        raw = Path(path).read_bytes()
        if detect_binary_text(raw):
            raise ValueError("文件内容为二进制/压缩数据，与文本扩展名不符（请解压出文本文件或转换格式后再上传）")
        return read_text_any_encoding(raw)
    if ext == ".docx":
        from docx import Document

        doc = Document(path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for t in doc.tables:
            for row in t.rows:
                parts.append("| " + " | ".join(c.text.strip() for c in row.cells) + " |")
        return "\n".join(parts)
    if ext == ".pptx":
        from pptx import Presentation

        prs = Presentation(path)
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    parts.append(shape.text_frame.text)
        return "\n".join(parts)
    if ext == ".pdf":
        return _pdf_extract(path)[0]
    if ext == ".xlsx":
        return xlsx_preview(path)
    raise ValueError(f"不支持的格式 {ext}")


def _pdf_extract(path) -> tuple[str, int, int]:
    """PDF 单趟解析：返回 (文本, 嵌入图片数, 无文本页数)。

    2026-09-10：原先「parse_document 取文本 + _media_stats 再开一次数图片」= 同一 PDF 解析两遍，
    且 fitz 文档**从不 close** 导致内存跨文件累积（实测一个 22MB 的图文 PDF 把进程撑到 2.6GB RSS）。
    改为一次打开、一趟取完、finally 关闭。
    """
    import fitz

    pdf = fitz.open(path)
    try:
        parts: list[str] = []
        images = textless = 0
        for page in pdf:
            t = page.get_text()
            parts.append(t)
            images += len(page.get_images())
            if len(t.strip()) < 20:
                textless += 1
        return "\n".join(parts), images, textless
    finally:
        pdf.close()


# 兼容旧引用名（原 kb_service.parse_kb_document）
parse_kb_document = parse_document


# ---------- 格式嗅探 ----------

# 日期锚点：**允许日期行带装饰**（`## 2026-07-31` / `********* 2026-07-12 *********` /
# `=== 2026-07-12 ===` 等）——按"整行除日期外只有非单词字符"判定，不认死某一种导出格式。
_CHAT_DAY_RE = re.compile(r"^[^\w\n]{0,12}(\d{4}-\d{2}-\d{2})[^\w\n]{0,12}$", re.M)
# 消息行候选：(说话人捕获正则, 判定阈值)。弱模式（无装饰的 `名字: 内容`）阈值抬高防误伤——
# 阈值 8 是实测折中：真实聊天记录导出动辄数百行，而普通文档里「字段: 值」很少连出 8 行
# 又恰好有独立日期行的组合。误判代价也低（仍按行边界成块，只是块标题带上日期/参与者）。
_CHAT_MSG_RES = (
    (re.compile(r"^\*\*([^*\n]{1,30})\*\*\s*[:：]\s?", re.M), 5),    # md 导出：**发言人**: 内容
    (re.compile(r"^([^\s:：\n][^:：\n]{0,28})\s*[:：]\s", re.M), 8),  # txt 导出：发言人: 内容
)
_SNIFF_CHARS = 200_000   # 只看前 20 万字符（足够判定，避免超长文件全扫）


def detect_chat(text: str) -> re.Pattern | None:
    """聊天记录嗅探：命中返回「消息行正则」（说话人在第 1 组），否则 None。

    双条件：① 至少一个日期锚点行；② 某种消息行模式达到阈值。
    按**内容结构**判定，不看文件名/扩展名——换名字、换装饰符（md / txt 两种导出格式）都识别；
    不命中即回退通用策略，永不报错。
    """
    head = text[:_SNIFF_CHARS]
    if not _CHAT_DAY_RE.search(head):
        return None
    for rex, threshold in _CHAT_MSG_RES:
        if len(rex.findall(head)) >= threshold:
            return rex
    return None


def sniff_format(text: str) -> str:
    """结构嗅探：命中聊天记录导出特征 → "chat"；否则 "plain"（永不报错）。"""
    return "chat" if detect_chat(text) else "plain"


# ---------- 聊天记录切块 ----------

_PUNCT = "。！？；.!?;"


def _chat_units(text: str, msg_re: re.Pattern) -> list[dict]:
    """把聊天记录拆成「消息单元」：{day, speaker, text, start}。

    日期行只负责切换当前日期（该行本身不进正文）；每条消息起一条单元，直到下一条
    消息/下个日期为止（XML 系统消息、合并转发那样的嵌套列表自然并进所属消息）。
    """
    marks: list[tuple[int, str, str]] = []
    for m in _CHAT_DAY_RE.finditer(text):
        marks.append((m.start(), "day", m.group(1)))
    for m in msg_re.finditer(text):
        marks.append((m.start(), "msg", m.group(1)))
    marks.sort()
    units: list[dict] = []
    cur_day = ""
    for i, (pos, kind, val) in enumerate(marks):
        end = marks[i + 1][0] if i + 1 < len(marks) else len(text)
        if kind == "day":
            cur_day = val
            rest = text[pos:end]
            nl = rest.find("\n")
            body = rest[nl + 1:] if nl >= 0 else ""
            if body.strip():
                off = pos + nl + 1 + (len(body) - len(body.lstrip()))
                units.append({"day": cur_day, "speaker": "", "text": body.strip(), "start": off})
            continue
        seg = text[pos:end].strip()
        if seg:
            units.append({"day": cur_day, "speaker": val, "text": seg, "start": pos})
    if marks and marks[0][0] > 0:      # 首个标记之前的前言（如导出说明）
        head = text[:marks[0][0]].strip()
        if head:
            units.insert(0, {"day": "", "speaker": "", "text": head, "start": 0})
    return units


def _split_long_message(unit: dict, max_size: int, msg_re: re.Pattern) -> list[dict]:
    """单条消息超长：按标点硬切，**每片补回发言人前缀**（不丢归属）。"""
    raw = unit["text"]
    m = msg_re.match(raw)
    head = f"**{unit['speaker']}**: " if unit["speaker"] else ""
    body = raw[m.end():] if m else raw
    budget = max(200, max_size - len(head))
    pieces: list[dict] = []
    base = 0
    while len(body) > budget:
        cut = body[:budget]
        idx = max(cut.rfind(c) for c in _PUNCT)
        idx = budget if idx < budget * 0.6 else idx + 1
        pieces.append({"day": unit["day"], "speaker": unit["speaker"],
                       "text": head + body[:idx], "start": unit["start"] + base})
        base += idx
        body = body[idx:]
    if body:
        pieces.append({"day": unit["day"], "speaker": unit["speaker"],
                       "text": head + body, "start": unit["start"] + base})
    return pieces


def _render_chat_block(units: list[dict]) -> tuple[str, str]:
    """块 → (块标题, 正文)。标题与正文首行都带「日期 · 参与者」，纯看正文也能定位。

    参与者按**块内出现次数**取前 6（说话人多的群聊里，出现一次的字段名如「链接: 」「获赞与收藏: 」
    会被真正的发言人挤出去——weak 模式（`名字: 内容`）的固有噪声，按频次过滤最省事）。
    """
    days = [u["day"] for u in units if u["day"]]
    span = ""
    if days:
        span = days[0] if days[0] == days[-1] else f"{days[0]} ~ {days[-1]}"
    freq: dict[str, int] = {}
    for u in units:
        if u["speaker"]:
            freq[u["speaker"]] = freq.get(u["speaker"], 0) + 1
    ordered = sorted(freq, key=lambda s: -freq[s])[:6]   # 稳定排序：同频次按首次出现先后
    head = " · ".join(x for x in (span, "/".join(ordered) + ("…" if len(freq) > 6 else "")) if x)
    title = head[:200]
    body = "\n".join(u["text"] for u in units)
    return title, (f"[{head}]\n{body}" if head else body)


def chunk_chat(text: str, msg_re: re.Pattern, min_size: int = MIN_SIZE,
               max_size: int = MAX_SIZE) -> list[dict]:
    """聊天记录切块：日期 / 消息双锚点，绝不从消息中间切开。

    - 一块尽量不跨天（换日且已够 min_size 即断块）
    - 重叠 = **回带上块最后 1 条完整消息**（不是 N 个字符——字符重叠会把消息劈两半）
    - 群聊记录量大时按 max_size 在消息边界断块
    - msg_re 由 detect_chat 嗅探得出（适配 md / txt 两种导出格式）
    """
    units = _chat_units(text, msg_re)
    if not units:
        return []
    blocks: list[dict] = []
    cur: list[dict] = []
    cur_len = 0
    msg_no = 0

    def flush() -> None:
        nonlocal cur, cur_len
        if not cur:
            return
        title, content = _render_chat_block(cur)
        blocks.append({
            "seq": len(blocks) + 1,
            "block_title": title,
            "content": content,
            "para_loc": f"msg {cur[0]['_no']}-{cur[-1]['_no']}",
            "char_start": cur[0]["start"],
        })
        tail = cur[-1]
        # 回带最后一条完整消息（过大则不回带，防下一块一开头就超限）
        carry = [tail] if len(tail["text"]) <= max_size // 2 else []
        cur, cur_len = carry, sum(len(u["text"]) for u in carry)

    for u in units:
        pieces = _split_long_message(u, max_size, msg_re) if len(u["text"]) > max_size else [u]
        for p in pieces:
            msg_no += 1
            p["_no"] = msg_no
            # 换日 / 超限 → 断块（flush 会把上块末条消息留作 carry，再追加本片）
            if cur and cur_len >= min_size:
                day_break = bool(p["day"] and cur[-1]["day"] and p["day"] != cur[-1]["day"])
                if day_break or cur_len + len(p["text"]) > max_size:
                    flush()
            cur.append(p)
            cur_len += len(p["text"])
    flush()
    return blocks


# ---------- 统一入口 ----------


def chunk_document(
    text: str,
    min_size: int = MIN_SIZE,
    max_size: int = MAX_SIZE,
    overlap: int = OVERLAP,
) -> list[dict]:
    """切块统一入口：嗅探格式 → 聊天策略 / 通用策略；异常一律回退通用策略。

    返回 [{seq, block_title, content, para_loc, char_start}]（与 kb_chunker.chunk_text 同构）。
    """
    if not text or not text.strip():
        return []
    try:
        msg_re = detect_chat(text)
        if msg_re is not None:
            blocks = chunk_chat(text, msg_re, min_size=min_size, max_size=max_size)
            if blocks:
                return blocks
    except Exception as e:
        logger.warning("聊天记录切块失败，回退通用策略: %s", str(e)[:150])
    return chunk_text(text, min_size=min_size, max_size=max_size, overlap=overlap)


# ---------- 检索管线（候选预筛 → LLM 精排） ----------

# 中英同义扩展（英文文档对中文关键词零命中的轻量缓解）
_SYNONYMS = {
    "文档": ["document", "doc", "md"], "知识": ["knowledge", "kb", "wiki"], "规范": ["spec", "standard", "guideline"],
    "流程": ["process", "workflow", "flow"], "制度": ["policy", "regulation"], "测试": ["test", "testing", "qa"],
    "报告": ["report", "summary"], "数据": ["data", "dataset"], "统计": ["stats", "metrics", "summary"],
    "营销": ["marketing"], "表格": ["table", "sheet"], "模板": ["template"], "示例": ["example", "sample"],
    "指南": ["guide", "manual", "howto"], "说明": ["readme", "note", "description"], "前端": ["frontend", "front-end", "ui"],
    "后端": ["backend", "server", "api"], "工程": ["engineering", "dev"], "架构": ["architecture", "arch"],
    "合同": ["contract", "agreement"], "客户": ["customer", "client"], "聊天": ["chat", "message", "history"],
    "直播": ["live", "livestream"], "素材": ["asset", "material"], "报价": ["quote", "price"],
}


def expand_synonyms(keywords: list[str]) -> list[str]:
    """关键词中英同义扩展（去重，≤12 个防 SQL 膨胀）。"""
    out = list(keywords)
    for kw in keywords:
        for syn in _SYNONYMS.get(kw, []):
            if syn not in out:
                out.append(syn)
    return out[:12]


async def extract_keywords(query: str, cfg: dict | None) -> list[str]:
    """LLM 提检索关键词（1-3 个实质词）；失败/超时回退整句。走 llm_aux 轻量档（cfg 由调用方解析）。"""
    import json as _json

    from app.agent.llm_client import create_text_client, json_mode_kwargs

    try:
        client, model, kwargs = create_text_client(cfg)
        kwargs = {**kwargs, **json_mode_kwargs(cfg)}   # GLM 结构化输出（2026-09-11）
        resp = await client.chat.completions.create(
            model=model,
            messages=[
                {"role": "system",
                 "content": "从用户问题中提取 1-3 个用于内部资料检索的关键词。"
                            "只输出实质性词语（名词、主题、专有名词、人名、品牌名），不要疑问词（什么/怎么/哪里/是否/为什么）和虚词（的/了/呢/吗/请/帮我/可以）。"
                            "输出 JSON 字符串数组，如 [\"周报\", \"提交\"]。只输出 JSON。"},
                {"role": "user", "content": query},
            ],
            **kwargs,
        )
        kws = _json.loads((resp.choices[0].message.content or "[]").strip().strip("```json").strip("```"))
        cleaned = [str(k).strip() for k in kws if isinstance(k, (str, int)) and str(k).strip()]
        if cleaned:
            return cleaned[:3]
    except Exception:
        pass
    return [query]


async def rank_candidates(query: str, candidates: list, top_k: int, cfg: dict | None,
                          render) -> list:
    """LLM 精排：从候选中选 top_k（render(cand) 产出给 LLM 的一行摘要）；失败回退前 top_k。"""
    import json as _json

    if len(candidates) <= top_k:
        return list(candidates[:top_k])
    listing = "\n".join(f"[{i}] {render(c)}" for i, c in enumerate(candidates))
    from app.agent.llm_client import create_text_client, json_mode_kwargs

    try:
        client, model, kwargs = create_text_client(cfg)
        kwargs = {**kwargs, **json_mode_kwargs(cfg)}   # GLM 结构化输出（2026-09-11）
        resp = await client.chat.completions.create(
            model=model,
            messages=[
                {"role": "system",
                 "content": f"从候选资料中选出与用户问题最相关的 {top_k} 条，输出 JSON 数组（序号，按相关度降序），如 [0, 3]。只输出 JSON。"},
                {"role": "user", "content": f"问题：{query}\n候选：\n{listing}"},
            ],
            **kwargs,
        )
        idxs = _json.loads((resp.choices[0].message.content or "[]").strip().strip("```json").strip("```"))
        picked: list = []
        if isinstance(idxs, list):
            for i in idxs:
                if isinstance(i, int) and 0 <= i < len(candidates):
                    picked.append(candidates[i])
                if len(picked) >= top_k:
                    break
        if picked:
            return picked
    except Exception:
        pass
    return list(candidates[:top_k])


# ---------- 解析分发（L1~L4） ----------


@dataclass
class ExtractResult:
    """解析结果 + 文件画像（画像随块一起返回，防 agent 误判"搜不到就是没有"）。"""
    ok: bool
    text: str = ""
    quality: str = "high"          # high（专用解析/LibreOffice）| low（粗暴捞取）
    kind: str = "text"             # text | hybrid | asset | unreadable
    reason: str = ""
    images: int = 0                # 嵌入图片数（docx/pptx/pdf）
    textless_pages: int = 0        # 无文本层的页数（PDF 扫描页）
    # 2026-09-18：xlsx/xlsm 逐 sheet 预览 [(sheet 名, 文本)]——供「一 sheet 一块」入库；
    # text 字段仍是拼接后整体截断的老口径（只吃字符串的调用方行为不变）
    sheets: list[tuple[str, str]] = field(default_factory=list)

    def profile(self) -> dict:
        return {"kind": self.kind, "text_chars": len(self.text),
                "images": self.images, "textless_pages": self.textless_pages}


_LO_TIMEOUT_S = 150


async def _libreoffice_extract(path: Path) -> str | None:
    """L2：用 LibreOffice 把老 Office 家族转 UTF-8 文本（**复用 services/libreoffice**）。

    实测 4.3MB .doc → 13.7k 字干净正文 / 34s（首次含建 profile），故仅用于**后台索引**；
    单实例信号量 + 独立 profile + 可配置 soffice 路径都由 services/libreoffice 负责
    （2026-09-10 同类问题排查：原实现自建 subprocess 且硬编码 "soffice"，违反路径配置约定）。
    失败/超时返回 None 交 L3 兜底。
    """
    from app.services import libreoffice as lo

    out_dir = Path(tempfile.mkdtemp(prefix="docidx_lo_"))
    try:
        try:
            await lo.run_convert(
                ["--convert-to", "txt:Text (encoded):UTF8", "--outdir", str(out_dir), str(path)],
                timeout=_LO_TIMEOUT_S)
        except Exception as e:
            logger.warning("LibreOffice 转换失败 %s: %s", path.name, str(e)[:120])
            return None
        for f in out_dir.glob("*.txt"):
            return f.read_text(encoding="utf-8", errors="ignore")
        return None
    finally:
        await asyncio.to_thread(shutil.rmtree, out_dir, True)


# L3 长串过滤：老 Office 的 OLE 流里夹带格式/图片字节，只留"像句子"的片段
_RUN_RE = re.compile(
    r"[一-鿿][一-鿿　-〿！-～，。、；：？！“”‘’（）《》]{7,}"
    r"|[A-Za-z0-9][A-Za-z0-9 ,.\-_:;()/]{19,}"
)


def _naive_extract(raw: bytes) -> str:
    """L3：未知二进制的粗暴文本捞取（三种解码试探，取捞得最多的一种）。"""
    best = ""
    for enc in ("utf-16-le", "utf-8", "gbk"):
        try:
            text = raw.decode(enc, errors="ignore")
        except Exception:
            continue
        runs = _RUN_RE.findall(text)
        if not runs:
            continue
        out = "\n".join(runs)
        if len(out) > len(best):
            best = out
    return best


def _pdf_textless_pages(path: str) -> int:
    """PDF 无文本层页数（扫描页）——只判每页文本长度，轻量且显式关闭文档。"""
    try:
        import fitz

        pdf = fitz.open(path)
        try:
            return sum(1 for page in pdf if len(page.get_text().strip()) < 20)
        finally:
            pdf.close()
    except Exception:
        return 0


async def extract_text(path: str) -> ExtractResult:
    """解析分发：按内容与扩展名逐层兜底，任何情况都返回结果（不抛异常）。"""
    p = Path(path)
    ext = p.suffix.lower()
    if not p.exists():
        return ExtractResult(ok=False, kind="unreadable", reason="文件不存在")

    # L1 专用解析器
    try:
        if ext in {".xlsx", ".xlsm"}:
            sheets = await asyncio.to_thread(xlsx_preview_sheets, str(p))
            # text 保持老口径（拼接后整体截断）——只吃字符串的调用方行为不变；
            # sheets 另供「一 sheet 一块」入库（多 sheet 不再被整体截断丢尾）
            text = "\n\n".join(t for _s, t in sheets)[:XLSX_PREVIEW_MAX_CHARS].strip()
            return ExtractResult(ok=bool(text), text=text,
                                 sheets=sheets if text else [],
                                 kind="text" if text else "unreadable",
                                 reason="" if text else "表格无数据")
        if ext == ".pdf":
            # 文本走 markitdown（与 file_parse 同源）；无文本页另用 fitz 轻量数一遍（只判长度）
            text = await asyncio.to_thread(parse_document, str(p), ext)
            textless = await asyncio.to_thread(_pdf_textless_pages, str(p))
        elif ext in TEXT_EXTS or ext in {".docx", ".pptx"}:
            text = await asyncio.to_thread(parse_document, str(p), ext)
            textless = 0
        else:
            text = None
        if text is not None:
            images = text.count("![")   # markitdown 的图片占位符 = 该文件嵌入的图片（与索引文本同源）
            text = text.strip()
            if not text:
                kind = "asset" if ext in ASSET_EXTS else "unreadable"
                reason = "扫描件/无文本层" if ext == ".pdf" and textless else "无可提取文本"
                return ExtractResult(ok=False, kind=kind, reason=reason,
                                     images=images, textless_pages=textless)
            kind = "hybrid" if (images or textless) else "text"
            return ExtractResult(ok=True, text=text, kind=kind,
                                 images=images, textless_pages=textless)
    except Exception as e:
        logger.warning("专用解析失败（转下层兜底）%s: %s", p.name, str(e)[:150])

    # L2 老 Office 家族（LibreOffice）
    if ext in OFFICE_OLD_EXTS:
        text = await _libreoffice_extract(p)
        if text and text.strip():
            return ExtractResult(ok=True, text=text.strip(), kind="text")
    elif ext in ASSET_EXTS:
        return ExtractResult(ok=False, kind="asset", reason="纯资源文件（无文本）")

    # L3 粗暴捞取（只对非资源类二进制）
    if ext not in ASSET_EXTS:
        try:
            raw = await asyncio.to_thread(p.read_bytes)
            text = _naive_extract(raw[:20 * 1024 * 1024])   # 最多看前 20MB
            if text.strip():
                return ExtractResult(ok=True, text=text.strip(), quality="low", kind="hybrid",
                                     reason="无专用解析器，文本为粗暴提取（可能有噪声）")
        except Exception as e:
            logger.warning("粗暴提取失败 %s: %s", p.name, str(e)[:120])

    # L4 取不到
    return ExtractResult(ok=False, kind="unreadable", reason=f"暂不支持解析该格式（{ext or '无扩展名'}）")
